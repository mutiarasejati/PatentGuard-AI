"""
PatentGuard AI - Exportable FTO (Freedom to Operate) Report Generator
Menghasilkan laporan resmi format PDF, Word (.docx), dan HTML siap cetak.
Memuat Ringkasan Eksekutif, Dossier Kimia & IUPAC, Matriks Klaim Farmasi, dan Rekomendasi Reformulasi.
"""

import io
import datetime
from typing import Optional
from models import PharmaInfringementReport, InfringementRiskLevel


class FTOReportGenerator:
    """Mesin pembuatan dokumen laporan FTO profesional."""

    @staticmethod
    def generate_pdf(report: PharmaInfringementReport) -> bytes:
        """Menghasilkan file PDF biner menggunakan ReportLab."""
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import (
            SimpleDocTemplate,
            Paragraph,
            Spacer,
            Table,
            TableStyle,
            HRFlowable,
            KeepTogether,
        )

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Heading1"],
            fontSize=20,
            leading=24,
            textColor=colors.HexColor("#1e1b4b"),
            spaceAfter=4,
        )
        subtitle_style = ParagraphStyle(
            "DocSubtitle",
            parent=styles["Normal"],
            fontSize=10,
            leading=13,
            textColor=colors.HexColor("#475569"),
            spaceAfter=12,
        )
        heading2_style = ParagraphStyle(
            "Heading2Custom",
            parent=styles["Heading2"],
            fontSize=13,
            leading=16,
            textColor=colors.HexColor("#312e81"),
            spaceBefore=12,
            spaceAfter=6,
        )
        body_style = ParagraphStyle(
            "BodyCustom",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#1e293b"),
        )
        table_header_style = ParagraphStyle(
            "TableHeader",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            fontName="Helvetica-Bold",
            textColor=colors.white,
        )
        table_cell_style = ParagraphStyle(
            "TableCell",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#0f172a"),
        )

        elements = []

        # Header Banner
        elements.append(Paragraph("PATENTGUARD AI - FREEDOM TO OPERATE (FTO) DOSSIER", title_style))
        elements.append(Paragraph(
            f"Laporan Analisis Hukum Paten Farmasi & Rekayasa Formulasi | Tanggal: {datetime.datetime.now().strftime('%d %B %Y')}",
            subtitle_style
        ))
        elements.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor("#4338ca"), spaceAfter=14))

        # Overview Table
        risk_color = colors.HexColor("#ef4444") if report.overall_risk == InfringementRiskLevel.HIGH else (
            colors.HexColor("#f59e0b") if report.overall_risk == InfringementRiskLevel.MEDIUM else colors.HexColor("#10b981")
        )

        overview_data = [
            [
                Paragraph("<b>Produk Sasaran:</b>", body_style),
                Paragraph(report.product_name, body_style),
                Paragraph("<b>Tingkat Risiko:</b>", body_style),
                Paragraph(f"<font color='{risk_color.hexval()}'><b>{report.overall_risk.value} ({report.risk_score_numeric}%)</b></font>", body_style),
            ],
            [
                Paragraph("<b>Paten Pembanding:</b>", body_style),
                Paragraph(report.patent_title, body_style),
                Paragraph("<b>Bentuk Sediaan:</b>", body_style),
                Paragraph(report.dosage_form or "Tablet", body_style),
            ],
            [
                Paragraph("<b>All-Elements Rule:</b>", body_style),
                Paragraph("Terpenuhi (Literal Infringement)" if report.all_elements_rule_met else "Tidak Terpenuhi (Lolos)", body_style),
                Paragraph("<b>Doktrin Ekuivalensi:</b>", body_style),
                Paragraph("Terpicu (DoE Risk)" if report.doctrine_of_equivalents_triggered else "Tidak Terpicu", body_style),
            ],
        ]
        t_overview = Table(overview_data, colWidths=[110, 160, 110, 160])
        t_overview.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#cbd5e1")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        elements.append(t_overview)
        elements.append(Spacer(1, 12))

        # Executive Summary
        elements.append(Paragraph("1. RINGKASAN EKSEKUTIF & KESIMPULAN FTO", heading2_style))
        elements.append(Paragraph(f"<b>Opini FTO:</b> {report.fto_verdict}", body_style))
        elements.append(Spacer(1, 4))
        elements.append(Paragraph(report.executive_summary, body_style))
        elements.append(Spacer(1, 10))

        # Section 2: Chemical & IUPAC Dossier
        elements.append(Paragraph("2. VERIFIKASI TATA NAMA KIMIA & ZAT AKTIF (IUPAC DOSSIER)", heading2_style))
        if report.chemical_entities:
            chem_table_data = [
                [
                    Paragraph("Nama INN", table_header_style),
                    Paragraph("Nama IUPAC", table_header_style),
                    Paragraph("Bentuk / Garam", table_header_style),
                    Paragraph("Rumus Molekul", table_header_style),
                    Paragraph("Status Ekuivalensi", table_header_style),
                ]
            ]
            for c in report.chemical_entities:
                chem_table_data.append([
                    Paragraph(f"<b>{c.inn_name}</b><br/>({', '.join(c.trade_names[:2])})", table_cell_style),
                    Paragraph(c.iupac_name, table_cell_style),
                    Paragraph(c.chemical_form, table_cell_style),
                    Paragraph(f"{c.molecular_formula or '-'}<br/>CAS: {c.cas_number or '-'}", table_cell_style),
                    Paragraph("Identik / Ekuivalen" if c.is_equivalent_entity else "Zat Berbeda", table_cell_style),
                ])
            t_chem = Table(chem_table_data, colWidths=[95, 175, 95, 105, 70])
            t_chem.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#312e81")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            elements.append(t_chem)
        else:
            elements.append(Paragraph("Tidak ada zat aktif farmasi standar yang terdeteksi.", body_style))
        elements.append(Spacer(1, 10))

        # Section 3: Excipient Function & Equivalence Mapping
        elements.append(Paragraph("3. EVALUASI FUNGSI EKSIPIEN & FUNCTION-WAY-RESULT TEST", heading2_style))
        if report.excipient_evaluations:
            exc_table_data = [
                [
                    Paragraph("Eksipien Formula", table_header_style),
                    Paragraph("Peran Fungsional", table_header_style),
                    Paragraph("Komponen Paten", table_header_style),
                    Paragraph("Hasil Uji FWR (DoE)", table_header_style),
                    Paragraph("Status", table_header_style),
                ]
            ]
            for e in report.excipient_evaluations:
                fwr_text = e.fwr_test.legal_doctrine_summary if e.fwr_test else "Identik dengan klaim paten"
                exc_table_data.append([
                    Paragraph(f"<b>{e.ingredient_name}</b><br/>{e.unit_dosage_amount or ''}", table_cell_style),
                    Paragraph(e.role_description[:100] + "...", table_cell_style),
                    Paragraph(e.claimed_counterpart or "Tidak diklaim langsung", table_cell_style),
                    Paragraph(fwr_text[:110] + "...", table_cell_style),
                    Paragraph(e.equivalence_status.value, table_cell_style),
                ])
            t_exc = Table(exc_table_data, colWidths=[100, 140, 95, 130, 75])
            t_exc.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            elements.append(t_exc)
        elements.append(Spacer(1, 10))

        # Section 4: Claim Chart Matrix
        elements.append(Paragraph("4. MATRIKS CLAIM CHART ELEMEN-DEMI-ELEMEN", heading2_style))
        for cr in report.claim_results:
            elements.append(Paragraph(
                f"<b>Klaim {cr.claim_number} ({cr.claim_type.value})</b> - Risiko: <b>{cr.risk_level.value}</b>",
                body_style
            ))
            elements.append(Paragraph(f"<i>{cr.summary}</i>", body_style))
            elements.append(Spacer(1, 3))

            chart_data = [
                [
                    Paragraph("#", table_header_style),
                    Paragraph("Kategori", table_header_style),
                    Paragraph("Limitasi Klaim Paten", table_header_style),
                    Paragraph("Fitur Formula Produk", table_header_style),
                    Paragraph("Status", table_header_style),
                ]
            ]
            for elem in cr.elements:
                chart_data.append([
                    Paragraph(str(elem.element_index), table_cell_style),
                    Paragraph(elem.category.value, table_cell_style),
                    Paragraph(elem.claim_limitation, table_cell_style),
                    Paragraph(elem.formula_feature or "-", table_cell_style),
                    Paragraph(elem.status.value, table_cell_style),
                ])
            t_chart = Table(chart_data, colWidths=[20, 85, 175, 175, 85])
            t_chart.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#475569")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            elements.append(t_chart)
            elements.append(Spacer(1, 8))

        # Section 5: Design-Around Recommendations
        if report.design_around_recommendations:
            elements.append(KeepTogether([
                Paragraph("5. REKOMENDASI DESIGN-AROUND & REFORMULASI TAKTIS", heading2_style),
                HRFlowable(width="100%", thickness=1, color=colors.HexColor("#f59e0b"), spaceAfter=6),
            ]))
            for idx, da in enumerate(report.design_around_recommendations, start=1):
                da_content = [
                    Paragraph(f"<b>Strategi {idx}: {da.title}</b> [{da.strategy_type}]", body_style),
                    Paragraph(f"<b>Deskripsi:</b> {da.description}", body_style),
                    Paragraph(f"<b>Dasar Ilmiah FTO:</b> {da.scientific_rationale}", body_style),
                    Spacer(1, 4),
                ]
                elements.append(KeepTogether(da_content))

        # Disclaimer Footer
        elements.append(Spacer(1, 14))
        elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#94a3b8"), spaceAfter=4))
        elements.append(Paragraph(
            "<font size='7' color='#64748b'><b>DISCLAIMER HUKUM HAKI & FARMASETIKA:</b> Laporan ini diproduksi oleh PatentGuard AI "
            "untuk panduan teknis pra-formulasi dan analisis kebebasan beroperasi (FTO). Bukan merupakan legal opinion resmi yang mengikat "
            "pengadilan. Harap lakukan konsultasi dengan Konsultan Paten Terdaftar dan uji stabilitas/disolusi farmasi laboratorium.</font>",
            body_style
        ))

        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_docx(report: PharmaInfringementReport) -> bytes:
        """Menghasilkan file Word (.docx) biner menggunakan python-docx."""
        import docx
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        doc = docx.Document()

        # Set Margins
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        # Title
        title_p = doc.add_paragraph()
        title_run = title_p.add_run("PATENTGUARD AI - FREEDOM TO OPERATE (FTO) DOSSIER")
        title_run.font.name = "Arial"
        title_run.font.size = Pt(16)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(30, 27, 75)

        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run(f"Laporan Analisis Hukum Paten Farmasi | Tanggal: {datetime.datetime.now().strftime('%d %B %Y')}")
        sub_run.font.size = Pt(9.5)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(100, 116, 139)

        # Overview Table
        t_meta = doc.add_table(rows=3, cols=2)
        t_meta.style = "Table Grid"
        rows = t_meta.rows
        rows[0].cells[0].text = f"Produk Sasaran: {report.product_name}"
        rows[0].cells[1].text = f"Tingkat Risiko: {report.overall_risk.value} ({report.risk_score_numeric}%)"
        rows[1].cells[0].text = f"Paten Pembanding: {report.patent_title}"
        rows[1].cells[1].text = f"Bentuk Sediaan: {report.dosage_form or 'Tablet'}"
        rows[2].cells[0].text = f"All-Elements Rule: {'Terpenuhi (Literal)' if report.all_elements_rule_met else 'Lolos Paten'}"
        rows[2].cells[1].text = f"Doctrine of Equivalents: {'Terpicu' if report.doctrine_of_equivalents_triggered else 'Tidak Terpicu'}"

        doc.add_paragraph()

        # Section 1
        h1 = doc.add_heading("1. Ringkasan Eksekutif & Opini FTO", level=1)
        h1.runs[0].font.color.rgb = RGBColor(49, 46, 129)
        p_fto = doc.add_paragraph()
        p_fto.add_run(f"Kesimpulan FTO: ").bold = True
        p_fto.add_run(report.fto_verdict)
        doc.add_paragraph(report.executive_summary)

        # Section 2
        h2 = doc.add_heading("2. Verifikasi Zat Aktif & Tata Nama Kimia IUPAC", level=1)
        h2.runs[0].font.color.rgb = RGBColor(49, 46, 129)
        if report.chemical_entities:
            t_chem = doc.add_table(rows=1, cols=4)
            t_chem.style = "Table Grid"
            hdr = t_chem.rows[0].cells
            hdr[0].text = "Nama INN / Dagang"
            hdr[1].text = "Nama IUPAC"
            hdr[2].text = "Bentuk / Garam"
            hdr[3].text = "Status Ekuivalensi"
            for c in report.chemical_entities:
                row = t_chem.add_row().cells
                row[0].text = f"{c.inn_name} ({', '.join(c.trade_names[:2])})"
                row[1].text = c.iupac_name
                row[2].text = c.chemical_form
                row[3].text = "Identik / Ekuivalen" if c.is_equivalent_entity else "Zat Berbeda"

        doc.add_paragraph()

        # Section 3
        h3 = doc.add_heading("3. Pemetaan Fungsi Eksipien & Uji FWR (Function-Way-Result)", level=1)
        h3.runs[0].font.color.rgb = RGBColor(49, 46, 129)
        if report.excipient_evaluations:
            t_exc = doc.add_table(rows=1, cols=4)
            t_exc.style = "Table Grid"
            hdr = t_exc.rows[0].cells
            hdr[0].text = "Eksipien Formula"
            hdr[1].text = "Peran Fungsional"
            hdr[2].text = "Uji FWR Doktrin Ekuivalensi"
            hdr[3].text = "Status"
            for e in report.excipient_evaluations:
                row = t_exc.add_row().cells
                row[0].text = f"{e.ingredient_name} ({e.unit_dosage_amount or ''})"
                row[1].text = e.role_description[:80] + "..."
                row[2].text = e.fwr_test.legal_doctrine_summary if e.fwr_test else "Identik dengan klaim paten"
                row[3].text = e.equivalence_status.value

        doc.add_paragraph()

        # Section 4: Claim Chart
        h4 = doc.add_heading("4. Matriks Claim Chart Elemen-demi-Elemen", level=1)
        h4.runs[0].font.color.rgb = RGBColor(49, 46, 129)
        for cr in report.claim_results:
            p_c = doc.add_paragraph()
            p_c.add_run(f"Klaim {cr.claim_number} ({cr.claim_type.value}) - Risiko: {cr.risk_level.value}").bold = True
            doc.add_paragraph(cr.summary)

            t_elem = doc.add_table(rows=1, cols=4)
            t_elem.style = "Table Grid"
            hdr = t_elem.rows[0].cells
            hdr[0].text = "#"
            hdr[1].text = "Kategori"
            hdr[2].text = "Limitasi Klaim Paten"
            hdr[3].text = "Fitur Formula Produk"
            for elem in cr.elements:
                row = t_elem.add_row().cells
                row[0].text = str(elem.element_index)
                row[1].text = elem.category.value
                row[2].text = elem.claim_limitation
                row[3].text = f"{elem.formula_feature or '-'} [{elem.status.value}]"
            doc.add_paragraph()

        # Section 5: Design-Around
        if report.design_around_recommendations:
            h5 = doc.add_heading("5. Rekomendasi Taktis Design-Around & Reformulasi", level=1)
            h5.runs[0].font.color.rgb = RGBColor(49, 46, 129)
            for idx, da in enumerate(report.design_around_recommendations, start=1):
                p_da = doc.add_paragraph()
                p_da.add_run(f"Strategi {idx}: {da.title} ({da.strategy_type})").bold = True
                doc.add_paragraph(f"Deskripsi: {da.description}")
                doc.add_paragraph(f"Dasar Ilmiah: {da.scientific_rationale}")

        # Disclaimer
        p_disc = doc.add_paragraph()
        p_disc.add_run("DISCLAIMER: Dokumen analisis FTO ini dihasilkan oleh PatentGuard AI untuk keperluan sains teknologi farmasi dan penilaian kebebasan berusaha awal. Konsultasikan dengan patent attorney sebelum rilis pasar.").italic = True

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
